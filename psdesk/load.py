"""Read a Salesforce tabular export from xlsx or csv."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from psdesk.columns import OPP_NAME, OPP_REQUIRED, QUOTE_REQUIRED

SPREADSHEET = {".xlsx", ".xlsm", ".csv"}
SKIP_NAMES = {"total", "grand total", "subtotal"}


class ReportError(Exception):
    """The file is not the export this scorecard knows how to read."""


@dataclass(frozen=True)
class Table:
    path: Path
    sheet: str
    headers: list[str]
    rows: list[dict[str, object]]


def load_table(path: Path) -> Table:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        sheet, headers, rows = _read_csv(path)
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        sheet, headers, rows = _read_xlsx(path)
    else:
        raise ReportError(f"{path.name} is not an xlsx or csv export")
    return Table(path=path, sheet=sheet, headers=headers, rows=rows)


def load_directory(folder: Path) -> tuple[Table, Table]:
    """Pick the closing report and the quote-line report out of a folder."""
    folder = Path(folder)
    files = sorted(
        p
        for p in folder.iterdir()
        if p.is_file()
        and p.suffix.lower() in SPREADSHEET
        and not p.name.startswith("~$")
    )
    if not files:
        raise ReportError(f"No xlsx or csv exports in {folder}")

    opps: list[Table] = []
    quotes: list[Table] = []
    unknown: list[str] = []
    for path in files:
        table = load_table(path)
        kind = _kind(table.headers)
        if kind == "opps":
            opps.append(table)
        elif kind == "quotes":
            quotes.append(table)
        else:
            unknown.append(path.name)

    problems: list[str] = []
    if len(opps) != 1:
        found = ", ".join(t.path.name for t in opps) or "none"
        problems.append(f"expected one closing-opportunity export, found {found}")
    if len(quotes) != 1:
        found = ", ".join(t.path.name for t in quotes) or "none"
        problems.append(f"expected one quote-line export, found {found}")
    if unknown:
        problems.append("ignored " + ", ".join(unknown))
    if problems and (len(opps) != 1 or len(quotes) != 1):
        raise ReportError("; ".join(problems))
    return opps[0], quotes[0]


def require_reports(opps: Table, quotes: Table) -> None:
    _require(opps, OPP_REQUIRED, "closing-opportunity export")
    _require(quotes, QUOTE_REQUIRED, "quote-line export")


def _kind(headers: list[str]) -> str | None:
    have = set(headers)
    if set(OPP_REQUIRED).issubset(have):
        return "opps"
    if set(QUOTE_REQUIRED).issubset(have):
        return "quotes"
    return None


def _require(table: Table, required: list[str], label: str) -> None:
    missing = [name for name in required if name not in table.headers]
    if missing:
        found = ", ".join(table.headers)
        raise ReportError(
            f"{table.path.name} is not a {label}. Missing: {', '.join(missing)}. "
            f"Found: {found}"
        )


def _read_csv(path: Path) -> tuple[str, list[str], list[dict[str, object]]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        rows = list(reader)
    header_at = _header_index(rows)
    if header_at is None:
        raise ReportError(f"{path.name} has no Opportunity Name column")
    headers = _clean_headers(rows[header_at])
    data = [_row_dict(headers, raw) for raw in rows[header_at + 1 :]]
    return path.stem, headers, [row for row in data if row is not None]


def _read_xlsx(path: Path) -> tuple[str, list[str], list[dict[str, object]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            header_at = _header_index(rows)
            if header_at is None:
                continue
            headers = _clean_headers(rows[header_at])
            data = [_row_dict(headers, raw) for raw in rows[header_at + 1 :]]
            return sheet_name, headers, [row for row in data if row is not None]
    finally:
        workbook.close()
    raise ReportError(f"{path.name} has no Opportunity Name column")


def _header_index(rows: list) -> int | None:
    for index, row in enumerate(rows[:15]):
        if row and any(_cell_text(cell) == OPP_NAME for cell in row):
            return index
    return None


def _clean_headers(row) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for cell in row:
        name = _cell_text(cell)
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        headers.append(name)
    return headers


def _row_dict(headers: list[str], raw) -> dict[str, object] | None:
    values = list(raw) if raw is not None else []
    if len(values) < len(headers):
        values = values + [None] * (len(headers) - len(values))
    record = {headers[i]: values[i] for i in range(len(headers))}
    name = _cell_text(record.get(OPP_NAME))
    if not name or name.casefold() in SKIP_NAMES:
        return None
    return record


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()
